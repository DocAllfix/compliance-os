"""
findings.py — formato unico dei rilievi per il banco di misura.

Due sorgenti:
  - result.json di un'esecuzione (tools/bench/run_pack.py);
  - l'export DOCX "Full Findings Detail" (rilievi della demo, per tarare il
    giudice sulle etichette del committente).

Ogni rilievo diventa:
  {"ref": "9001 §7.2", "norm": "9001", "clause": "7.2", "title": ...,
   "severity": "MINOR", "score": 40, "finding": ..., "action": ...,
   "evidence": ["file.xlsx", ...], "is_spof": bool}
"""
from __future__ import annotations

import json
import os
import re
from typing import Any, Dict, List, Optional

SEVERITIES = ("MAJOR", "MINOR", "OBSERVATION", "CONFORMANT")


def _norm_short(norm: str) -> str:
    m = re.search(r"(\d{4,5})", norm or "")
    return m.group(1) if m else (norm or "?")


def _clause_from_id(clause_id: str) -> str:
    # "9001_8_4" → "8.4";  "14001_9_1_2" → "9.1.2"
    parts = str(clause_id).split("_")
    return ".".join(parts[1:]) if len(parts) > 1 else str(clause_id)


def _ref(norm: str, clause: str) -> str:
    return f"{norm} §{clause}"


def from_result(result: Dict[str, Any]) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    for n in result.get("norms", []):
        norm = _norm_short(n.get("norm", ""))
        for cl in n.get("clauses", []):
            clause = _clause_from_id(cl.get("clause_id", ""))
            rem = cl.get("remediation") or {}
            action = " ".join(x for x in (rem.get("what"), rem.get("correction"),
                                          rem.get("corrective_action")) if x)
            out.append({
                "ref": _ref(norm, clause), "norm": norm, "clause": clause,
                "title": cl.get("clause_title", ""),
                "severity": cl.get("severity", ""), "score": cl.get("score"),
                "finding": cl.get("finding", ""), "action": action,
                "evidence": sorted({s.get("file_name", "") for s in cl.get("sources", []) if s.get("file_name")}),
                "is_spof": bool(cl.get("is_spof")),
            })
    return out


def from_findings_docx(path: str) -> List[Dict[str, Any]]:
    import docx  # python-docx
    paras = [p.text.strip() for p in docx.Document(path).paragraphs if p.text.strip()]
    out: List[Dict[str, Any]] = []
    norm: Optional[str] = None
    cur: Optional[Dict[str, Any]] = None
    for p in paras:
        if re.fullmatch(r"ISO \d{4,5}:\d{4}", p):
            norm = _norm_short(p)
            continue
        m = re.match(r"^(\d+(?:\.\d+)*)\s+—\s+(.*)$", p)
        if m and norm:
            cur = {"ref": _ref(norm, m.group(1)), "norm": norm, "clause": m.group(1),
                   "title": m.group(2), "severity": "", "score": None, "finding": "",
                   "action": "", "evidence": [], "is_spof": False}
            out.append(cur)
            continue
        if cur is None:
            continue
        m = re.match(r"^Severity:\s*([A-Z]+)(\s*\[SPoF\])?.*?Score:\s*(\d+)", p)
        if m:
            cur["severity"], cur["is_spof"], cur["score"] = m.group(1), bool(m.group(2)), int(m.group(3))
        elif p.startswith("Finding:"):
            cur["finding"] = p[len("Finding:"):].strip()
        elif p.startswith("Evidence:"):
            cur["evidence"] = [e.strip() for e in p[len("Evidence:"):].split(";") if e.strip()]
        elif p.startswith("Recommended action:"):
            cur["action"] = p[len("Recommended action:"):].strip()
    return out


def extra_from_result(result: Dict[str, Any]) -> str:
    """Le sezioni del report oltre ai rilievi per clausola: un'ENC può emergere
    lì (scadenze, coerenza tra clausole, priorità) e l'auditor la conta."""
    out: List[str] = []
    es = result.get("executive_summary") or {}
    if es.get("top_priority_integrations"):
        out.append("[Top priorities] " + " | ".join(es["top_priority_integrations"]))
    if es.get("executive_narrative"):
        out.append("[Executive summary] " + es["executive_narrative"])
    er = result.get("expiry_report") or {}
    for d in (er.get("expired_documents") or []) + (er.get("expiring_soon") or []):
        out.append(f"[Expiry] {d.get('file_name')}: expiry {d.get('expiry_date')}, "
                   f"days overdue {d.get('days_overdue')}, {d.get('severity')}")
    for c in result.get("cross_clause_findings") or []:
        out.append(f"[Cross-clause {c.get('norm')} {c.get('title')}] {c.get('status')}: {c.get('narrative')}")
    for c in result.get("self_checks") or []:     # Piano motore Fasi 4-5 (report §6G, UI)
        if c.get("status") in ("INCONSISTENT", "MISSING", "TO_VERIFY"):
            cites = "; ".join(f"{os.path.basename(str(x.get('file_name') or ''))}: {x.get('quote') or ''}"[:200]
                              for x in c.get("cites") or [])
            out.append(f"[Own records check] {c['status']}: {c.get('description')} ({cites})")
    di = result.get("documented_info_report") or {}
    for it in di.get("items") or []:
        if it.get("status") == "ABSENT":
            out.append(f"[Documented information {it.get('norm')} {it.get('clause_id')}] ABSENT: {it.get('item')}")
    return "\n".join(out)


def extra_from_report_docx(path: str) -> str:
    """Testo del report DOCX completo dal sommario esecutivo al piano d'azione
    (esclusa la matrice delle clausole, già coperta dall'export dei rilievi)."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    lines: List[str] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            t = Paragraph(child, d).text.strip()
            if t:
                lines.append(t)
        elif tag == "tbl":
            for r in Table(child, d).rows:
                cells: List[str] = []
                for c in r.cells:
                    ct = c.text.strip()
                    if ct and (not cells or cells[-1] != ct):
                        cells.append(ct)
                if cells:
                    lines.append(" | ".join(cells))
    heads = [i for i, l in enumerate(lines) if re.match(r"^§\d+[A-Z]?\s+[A-Z]", l) and l.upper() == l]
    start = next((i for i in heads if "EXECUTIVE SUMMARY" in lines[i]), 0)
    end = next((i for i in heads if "CLAUSE MATRIX" in lines[i]), len(lines))
    return "\n".join(lines[start:end])


def load_run(path: str, extra_docx: Optional[str] = None) -> Dict[str, Any]:
    """Cartella di un'esecuzione, result.json, oppure export DOCX dei rilievi
    (con, facoltativo, il report DOCX completo per le altre sezioni).
    Ritorna {"label", "findings", "extra", "key_roles"}."""
    if path.lower().endswith(".docx"):
        return {"label": os.path.basename(path), "findings": from_findings_docx(path),
                "extra": extra_from_report_docx(extra_docx) if extra_docx else "", "key_roles": None}
    rpath = os.path.join(path, "result.json") if os.path.isdir(path) else path
    with open(rpath, encoding="utf-8") as f:
        result = json.load(f)
    snap = ((result.get("executive_summary") or {}).get("company_snapshot") or {})
    return {"label": os.path.basename(os.path.dirname(rpath)) if os.path.isdir(path) else os.path.basename(path),
            "findings": from_result(result), "extra": extra_from_result(result),
            "key_roles": snap.get("key_roles"), "self_checks": result.get("self_checks") or []}
