"""
synthetic_pack.py — pacco sintetico per il banco di misura (Piano motore, Fase 0).

Un sistema di gestione integrato 9001/14001/45001 di un'azienda inglese
inventata (Northwind Fabrication Ltd, Leeds) con:
  - non conformità innestate di tutte le classi (A norma↔documento, B qualità
    del singolo documento, C contraddizione tra documenti, D "dichiara ma non
    fa", E aspettativa di settore) — la verità di riferimento è in key.json,
    nello stesso formato della chiave del banco Box;
  - le trappole di lettura che hanno prodotto false assenze sul campo: Excel
    multi-foglio con celle vuote, organigramma come immagine dentro un Word,
    planimetria come immagine, certificati PDF di fornitori con nomi di persone
    esterne;
  - "accuse vietate": cose che ESISTONO nel pacco e che il motore non deve
    dichiarare mancanti, più obblighi italiani che non si applicano a un'azienda
    UK. Sul pacco sintetico il valore atteso di ogni accusa vietata è zero.

Tutto è inventato: nessun contenuto di clienti. Uso:
  python tools/bench/synthetic_pack.py --out <cartella>   → pack.zip + key.json
"""
from __future__ import annotations

import argparse
import io
import json
import os
import zipfile
from typing import Dict, List

ROOT = "Northwind IMS"
COMPANY = "Northwind Fabrication Ltd"


# ── costruttori di file ──────────────────────────────────────────────────────

def _docx(paragraphs: List[str], tables: List[List[List[str]]] = (), image_png: bytes = None,
          image_caption: str = "") -> bytes:
    import docx
    from docx.shared import Inches
    d = docx.Document()
    for p in paragraphs:
        if p.startswith("# "):
            d.add_heading(p[2:], level=1)
        elif p.startswith("## "):
            d.add_heading(p[3:], level=2)
        elif p == "[IMAGE]" and image_png:
            if image_caption:
                d.add_paragraph(image_caption)
            d.add_picture(io.BytesIO(image_png), width=Inches(6))
        else:
            d.add_paragraph(p)
    for rows in tables:
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx(sheets: Dict[str, List[List]]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(lines: List[str]) -> bytes:
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((60, y), line, fontsize=11)
        y += 18
    return doc.tobytes()


def _png_boxes(title: str, boxes: List[str], cols: int = 3) -> bytes:
    """Diagramma a riquadri con testo (organigramma, planimetria): un'immagine
    vera, leggibile solo via OCR."""
    from PIL import Image, ImageDraw, ImageFont
    w, h = 1200, 220 + 140 * ((len(boxes) + cols - 1) // cols)
    img = Image.new("RGB", (w, h), "white")
    dr = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 26)
        tfont = ImageFont.truetype("arialbd.ttf", 34)
    except OSError:
        font = tfont = ImageFont.load_default()
    dr.text((40, 30), title, fill="black", font=tfont)
    bw = (w - 80 - (cols - 1) * 30) // cols
    for i, label in enumerate(boxes):
        r, c = divmod(i, cols)
        x0, y0 = 40 + c * (bw + 30), 110 + r * 140
        dr.rectangle([x0, y0, x0 + bw, y0 + 100], outline="black", width=3)
        for k, part in enumerate(label.split("\n")):
            dr.text((x0 + 14, y0 + 14 + k * 32), part, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# ── contenuti ────────────────────────────────────────────────────────────────

def build_files() -> Dict[str, bytes]:
    f: Dict[str, bytes] = {}

    org_chart = _png_boxes("Northwind Fabrication Ltd - Organisation Chart", [
        "Managing Director\nLaura Pennington", "QHSE Manager\nTom Ashworth",
        "Production Manager\nPriya Shah", "Workshop Supervisor\nDan Morley",
        "Purchasing Officer\nAmir Haddad", "Welders and Fabricators\n(6 operatives)",
    ])
    f["5 Leadership/IMS1 Integrated Management System Manual.docx"] = _docx([
        "# Integrated Management System Manual (IMS1) - Revision 3",
        "Northwind Fabrication Ltd designs and fabricates structural steel components at its single site in "
        "Leeds, United Kingdom. This manual describes the integrated management system for quality (ISO 9001:2015), "
        "environment (ISO 14001:2015) and occupational health and safety (ISO 45001:2018).",
        "## 1. Scope of the management system",
        "Scope: The design, cutting, welding, fabrication and painting of structural steel components for the "
        "construction sector, carried out at Unit 4, Aire Valley Works, Leeds LS10. The scope applies to all three "
        "standards. No requirement of ISO 9001 is excluded.",
        "## 2. Processes and their interactions",
        "The table below is the process map of the organisation (inputs, process, outputs, owner).",
        "## 3. Organisation and responsibilities",
        "The organisation chart below shows roles and reporting lines. Responsibilities and authorities for each "
        "role are described in the job descriptions held by the QHSE Manager.",
        "[IMAGE]",
        "## 4. Customer feedback",
        "All customer complaints and negative customer feedback are recorded in the NC Log (IMS12) and analysed at "
        "the management review.",
        "## 5. Context and interested parties",
        "Internal and external issues, interested parties and risks are maintained in the Context and Risk Register "
        "(IMS3).",
    ], tables=[[
        ["Process", "Inputs", "Outputs", "Owner"],
        ["Sales and contract review", "Customer enquiry, drawings", "Accepted order", "Managing Director"],
        ["Purchasing", "Material requisition", "Approved purchase order", "Purchasing Officer"],
        ["Cutting and fabrication", "Steel, drawings", "Fabricated components", "Production Manager"],
        ["Painting and dispatch", "Fabricated components", "Delivered order", "Workshop Supervisor"],
        ["Waste and environmental control", "Process waste, energy use", "Segregated waste, KPI data", "QHSE Manager"],
    ]], image_png=org_chart, image_caption="Figure 1 - Organisation chart")

    f["5 Leadership/IMS2 QHSE Policy.docx"] = _docx([
        "# QHSE Policy (IMS2)",
        "Northwind Fabrication Ltd is committed to satisfying customer requirements, preventing pollution, providing "
        "safe and healthy working conditions to prevent work-related injury and ill health, eliminating hazards and "
        "reducing OH&S risks, fulfilling its compliance obligations, consulting and involving workers, and "
        "continually improving the integrated management system.",
        "Signed: Laura Pennington, Managing Director. Date: 06/01/2026.",
    ])

    f["4 Context/IMS3 Context and Risk Register.xlsx"] = _xlsx({
        "Context": [["Issue", "Internal/External", "Effect on the IMS"],
                    ["Steel price volatility", "External", "Pressure on margins and supplier choice"],
                    ["Ageing overhead crane", "Internal", "Lifting risk and downtime"],
                    ["Local flood risk (River Aire)", "External", "Emergency planning, climate change relevant"]],
        "Interested Parties": [["Interested party", "Needs and expectations", "Compliance obligation?", "How monitored"],
                               ["Customers", "On-time delivery, CE marking to EN 1090", "Yes (contract)", "Delivery KPI"],
                               ["Employees", "Safe workplace, training", "Yes (law)", "H&S committee"],
                               ["Environment Agency", "Waste duty of care", "Yes (law)", "Legal register review"],
                               ["Neighbours", "Low noise in the evening", "No", "Complaints log"]],
        "Risks and Opportunities": [["Ref", "Risk / opportunity", "Linked issue or party", "Score", "Mitigation", "Owner"],
                                    ["R1", "Crane failure during lift", "Ageing overhead crane", 16, "Planned replacement 2027", "Production Manager"],
                                    ["R2", "Late steel deliveries", "Steel price volatility", 9, "Second approved stockholder", None],
                                    ["O1", "Solar panels on roof", "Environment Agency", 6, "Feasibility study", "QHSE Manager"]],
    })

    f["6 Planning/QHSE Objectives 2026.docx"] = _docx(["# QHSE Objectives 2026"], tables=[[
        ["Objective", "Target", "Owner", "Due", "Indicator"],
        ["Reduce rework on fabricated components", "Rework below 3% of hours", "Production Manager", "31/12/2026", "Rework hours / total hours"],
        ["Reduce electricity use per tonne", "-5% vs 2025", "QHSE Manager", "31/12/2026", "kWh per tonne"],
        ["Monthly safety inspection of the workshop", "12 inspections", "Workshop Supervisor", "", "Inspections completed"],
    ]])

    f["6 Planning/IMS4 Aspects and Impacts Register.xlsx"] = _xlsx({"Aspects": [
        ["Activity", "Aspect", "Impact", "Condition", "Significance (1-25)", "Control"],
        ["Welding", "Fume emissions", "Air quality", "Normal", 12, "Local exhaust ventilation"],
        ["Cutting", "Used cutting fluid", "Land / water contamination", "Normal", 15, "Hazardous waste segregation"],
        ["Painting", "Solvent VOCs", "Air quality", "Normal", 16, "Water-based paint trial"],
        ["Workshop", "Electricity use", "Resource depletion", "Normal", 10, "LED lighting"],
    ]})

    f["6 Planning/IMS7 Legal Register.xlsx"] = _xlsx({"Legal Register": [
        ["Legislation", "Requirement", "How we comply", "Evidence"],
        ["Health and Safety at Work etc. Act 1974", "General duties to employees", "IMS and risk assessments", "RA folder"],
        ["COSHH Regulations 2002", "Assess and control hazardous substances", "COSHH assessments and safety data sheets", "8 Operations/COSHH"],
        ["LOLER 1998", "Thorough examination of lifting equipment", "6-monthly examination of overhead crane", "8 Operations/Equipment"],
        ["Environmental Protection Act 1990 s.34", "Waste duty of care", "Licensed carriers, waste transfer notes", "8 Operations/Waste"],
        ["Hazardous Waste Regulations 2005", "Consignment notes for hazardous waste", "Consignment notes retained 3 years", "8 Operations/Waste"],
    ]})

    f["7 Support/IMS11 Training Matrix.xlsx"] = _xlsx({
        "Matrix": [["Name", "Role", "Manual Handling", "Abrasive Wheels", "Fire Warden", "Fire Warden trainer"],
                   ["Tom Ashworth", "QHSE Manager", "05/02/2026", "", "14/01/2026", "Tom Ashworth"],
                   ["Priya Shah", "Production Manager", "05/02/2026", "", "14/01/2026", "Tom Ashworth"],
                   ["Dan Morley", "Workshop Supervisor", "05/02/2026", "12/03/2026", "", ""],
                   ["Sam Okafor", "Welder", "05/02/2026", "12/03/2026", "", ""]],
        "Competence Requirements": [["Role", "Required training", "Refresher"],
                                    ["Workshop Supervisor", "Abrasive Wheels; Manual Handling; First Aid", "3 years"],
                                    ["Welder", "Abrasive Wheels; Manual Handling; Welding qualification EN ISO 9606", "2 years"],
                                    ["Production Manager", "Manual Handling; Fire Warden", "3 years"]],
    })
    f["7 Support/Training Records/Attendance - Manual Handling 05-02-2026.docx"] = _docx([
        "# Training Attendance Sheet", "Course: Manual Handling", "Date: 05/02/2026",
        "Trainer: ", "Attendees (signed): Tom Ashworth; Dan Morley; Sam Okafor",
    ])
    f["7 Support/Training Records/Attendance - Fire Warden 14-01-2026.docx"] = _docx([
        "# Training Attendance Sheet", "Course: Fire Warden", "Date: 14/01/2026",
        "Trainer: Gareth Lyle (Safe Skills Training Ltd)", "Attendees (signed): Tom Ashworth; Priya Shah",
    ])
    f["7 Support/Procedures/SOP1 Document Control Procedure.docx"] = _docx([
        "# SOP1 Document Control Procedure - Revision 2",
        "Documents are approved by the QHSE Manager before issue. Each controlled document carries a revision table "
        "at the end of the document recording the revision number, the change made and the approver.",
        "Obsolete documents are removed from the shared drive and archived in the Superseded folder.",
    ])

    f["7 Support/Calibration/IMS9 Calibration Register.xlsx"] = _xlsx({"Register": [
        ["ID", "Instrument", "Certificate", "Last calibration", "Result", "Next due"],
        ["CAL-01", "Torque wrench 40-200 Nm", "NWC-2201", "10/02/2026", "Pass", "10/09/2026"],
        ["CAL-02", "Pressure gauge 0-10 bar", "NWC-2202", "11/02/2026", "Pass", "11/08/2026"],
        ["CAL-03", "Platform scale 600 kg", "NWC-2203", "12/02/2026", "Pass", "12/08/2026"],
    ]})
    f["7 Support/Calibration/NWC-2201.pdf"] = _pdf([
        "Calibrex Services Ltd - Certificate of Calibration NWC-2201",
        "Instrument: Torque wrench 40-200 Nm (CAL-01)", "Date of calibration: 10/02/2026",
        "Result: PASS - within tolerance", "Next calibration due: 10/08/2026",
        "Calibrated by: Helen Marsh, Technical Manager, Calibrex Services Ltd",
    ])
    f["7 Support/Calibration/NWC-2202.pdf"] = _pdf([
        "Calibrex Services Ltd - Certificate of Calibration NWC-2202",
        "Instrument: Pressure gauge 0-10 bar (CAL-02)", "Date of calibration: 11/02/2026",
        "Result: FAIL - reading error +0.8 bar exceeds tolerance. Instrument must not be used.",
        "Calibrated by: Helen Marsh, Technical Manager, Calibrex Services Ltd",
    ])

    f["8 Operations/Procedures/SOP4 Purchasing and Supplier Control v2.docx"] = _docx([
        "# SOP4 Purchasing and Supplier Control", "Revision: 1.0",
        "Suppliers are approved by the Purchasing Officer before use. Every approved supplier must return the "
        "supplier questionnaire (TEMP4) and, where applicable, a copy of its certification; both are stored in "
        "8 Operations/Suppliers/<supplier name>. Approved suppliers are listed in IMS5 and re-evaluated yearly.",
    ])
    f["8 Operations/Suppliers/IMS5 Approved Supplier List.xlsx"] = _xlsx({"Approved suppliers": [
        ["Supplier", "Product / service", "Approval basis", "Approved on", "Re-evaluation due"],
        ["Calibrex Services Ltd", "Calibration", "ISO/IEC 17025 certificate; questionnaire", "15/01/2026", "15/01/2027"],
        ["Yorkshire Steel Stockholders", "Structural steel", "Questionnaire; mill certificates", "20/01/2026", "20/01/2027"],
        ["Greenline Waste Ltd", "Waste collection", "Waste carrier licence CBDU-44781", "22/01/2026", "22/01/2027"],
    ]})
    f["8 Operations/Suppliers/Calibrex Services Ltd/Supplier questionnaire.docx"] = _docx([
        "# Supplier questionnaire (TEMP4) - Calibrex Services Ltd",
        "Quality certification: ISO/IEC 17025 accredited (UKAS 0912).", "Completed by: Helen Marsh, Technical Manager.",
    ])
    f["8 Operations/Suppliers/Calibrex Services Ltd/UKAS certificate.pdf"] = _pdf([
        "UKAS - Certificate of Accreditation 0912", "Calibrex Services Ltd is accredited to ISO/IEC 17025:2017",
        "Signed for Calibrex: Helen Marsh, Technical Manager",
    ])
    f["8 Operations/Procedures/SOP5 Worker Consultation and Participation.docx"] = _docx([
        "# SOP5 Worker Consultation and Participation",
        "A health and safety committee with worker representatives meets quarterly. Committee minutes are stored in "
        "5 Leadership/Committee Minutes and actions are tracked to closure.",
        "Workers may raise hazards at any time through the supervisor or the suggestion box.",
    ])
    site_plan = _png_boxes("Site plan - fire extinguisher locations", [
        "FE1 CO2\nWelding bay", "FE2 CO2\nCutting area", "FE3 CO2\nPaint shop",
        "FE4 Foam\nPaint store", "FE5 CO2\nOffice", "FE6 CO2\nDispatch",
    ])
    f["8 Operations/Emergency/Site plan.png"] = site_plan
    f["8 Operations/Procedures/SOP6 Emergency Preparedness.docx"] = _docx([
        "# SOP6 Emergency Preparedness and Response",
        "There are seven fire extinguishers on site: five CO2 and two foam. Their locations are shown on the site plan.",
        "Fire drills are held every six months; the Fire Warden records the evacuation time.",
    ])
    f["8 Operations/Risk Assessments/RA1 Welding and Cutting.docx"] = _docx([
        "# RA1 Welding and Cutting - Risk Assessment",
        "Hazards: welding fume, UV radiation, burns, cutting fluid skin contact, lifting of steel sections with the "
        "5-tonne overhead crane.",
        "Controls: LEV extraction, welding screens, nitrile gloves, lifting plans, crane operated by trained staff only.",
        "Substances: argon/CO2 shielding gas, water-miscible cutting fluid, zinc-rich primer.",
        "Assessed by: Tom Ashworth, 12/01/2026.",
    ])

    f["9 Performance/IMS10 KPI Dashboard.xlsx"] = _xlsx({"KPIs 2026": [
        ["KPI", "Target", "Jan", "Feb", "Mar"],
        ["On-time delivery", "95%", "93%", "96%", "97%"],
        ["Rework hours %", "3%", "4.1%", "3.6%", "3.2%"],
        ["kWh per tonne", "-5%", "-1%", "-2%", "-4%"],
    ]})
    f["9 Performance/IMS8 Audit Programme 2026.xlsx"] = _xlsx({"Programme": [
        ["Audit", "Processes", "Planned date", "Auditor"],
        ["IA-01", "Purchasing, calibration", "14/04/2026", "Tom Ashworth"],
        ["IA-02", "Production, emergency", "15/09/2026", "External"],
    ]})
    f["9 Performance/Internal Audit Report IA-01 21-04-2026.docx"] = _docx([
        "# Internal Audit Report IA-01", "Date of audit: 21/04/2026", "Auditor: Tom Ashworth",
        "Processes audited: purchasing and calibration.",
        "Finding 1 (minor): purchase orders do not always reference the drawing revision.",
    ])
    f["9 Performance/Management Review Minutes 2026.docx"] = _docx([
        "# Management Review Minutes - 20/05/2026",
        "Attendees: Laura Pennington, Tom Ashworth, Priya Shah.",
        "Objectives: the rework objective (target 5% of hours) is on track. Electricity per tonne is down 4%.",
        "Internal audit IA-01 results reviewed; one minor finding closed.",
        "Environmental performance: aspects register reviewed, no new significant aspects.",
        "Resources: approval to replace the overhead crane in 2027.",
        "Actions: QHSE Manager to update the legal register by 30/06/2026.",
    ])

    f["10 Improvement/IMS12 NC Log.xlsx"] = _xlsx({"NC Log": [
        ["Ref", "Date raised", "Description", "Action", "Target date", "Status"],
        ["NC-013", "10/01/2026", "Wrong steel grade delivered", "Supplier informed, material returned", "31/01/2026", "Closed"],
        ["NC-014", "02/03/2026", "Weld undercut on beam batch 2207", "Retrain welder, add visual check", "31/03/2026", "Open"],
    ]})
    return f


# ── verità di riferimento ───────────────────────────────────────────────────

ENCS = [
    ("S01", "A", ["6.2"], "QHSE Objectives 2026", "The monthly safety inspection objective has no due date / time frame for completion."),
    ("S02", "A", ["9.1"], "IMS10 KPI Dashboard", "No occupational health and safety indicators are monitored (no OH&S KPIs or trend analysis in the KPI dashboard)."),
    ("S03", "A", ["9.3"], "Management Review Minutes 2026", "The management review does not discuss customer satisfaction or customer feedback."),
    ("S04", "A", ["6.1.2", "8.1"], "IMS4 Aspects and Impacts Register", "The aspects register does not consider the life cycle (no transport, use or end-of-life stages of the products)."),
    ("S05", "B", ["7.2", "7.5"], "Attendance - Manual Handling 05-02-2026", "The trainer field of the Manual Handling attendance sheet (05/02/2026) is blank."),
    ("S06", "B", ["10.2"], "IMS12 NC Log", "NC-014 is still open although its target date (31/03/2026) has passed."),
    ("S07", "B", ["7.5"], "SOP4 Purchasing and Supplier Control v2", "SOP4 file title says v2 but the document states Revision 1.0."),
    ("S08", "C", ["7.1.5", "9.1"], "IMS9 Calibration Register; NWC-2202", "The calibration register shows CAL-02 (certificate NWC-2202) as Pass, but the certificate says FAIL and the gauge must not be used."),
    ("S09", "C", ["6.2", "9.3"], "Management Review Minutes 2026; QHSE Objectives 2026", "The management review states a rework target of 5% while the objectives document sets 3%."),
    ("S10", "C", ["8.2"], "SOP6 Emergency Preparedness; Site plan", "SOP6 states seven extinguishers (five CO2, two foam) but the site plan shows six (five CO2, one foam)."),
    ("S11", "C", ["7.2"], "IMS11 Training Matrix; Attendance - Fire Warden 14-01-2026", "The training matrix records Tom Ashworth as the Fire Warden trainer, but the attendance sheet shows Gareth Lyle (Safe Skills Training Ltd)."),
    ("S12", "C", ["9.2"], "IMS8 Audit Programme 2026; Internal Audit Report IA-01", "Audit IA-01 was planned for 14/04/2026 but the report is dated 21/04/2026."),
    ("S13", "C", ["7.1.5", "9.1"], "IMS9 Calibration Register; NWC-2201", "For CAL-01 the register shows next due 10/09/2026 but certificate NWC-2201 says 10/08/2026."),
    ("S14", "D", ["7.2"], "IMS11 Training Matrix; Attendance - Manual Handling 05-02-2026", "Priya Shah is recorded as trained in Manual Handling on 05/02/2026 but she is not on the signed attendance sheet."),
    ("S15", "D", ["7.2"], "IMS11 Training Matrix", "Abrasive Wheels training on 12/03/2026 (Dan Morley, Sam Okafor) is on the matrix but there is no attendance record."),
    ("S16", "D", ["7.1.5", "9.1"], "IMS9 Calibration Register", "Calibration certificate NWC-2203 (CAL-03 platform scale) is listed in the register but missing from the pack."),
    ("S17", "D", ["8.4", "8.1.4"], "IMS5 Approved Supplier List; SOP4", "SOP4 requires a supplier questionnaire for every approved supplier, but there is none for Yorkshire Steel Stockholders (nor a licence copy for Greenline Waste)."),
    ("S18", "D", ["5.4"], "SOP5 Worker Consultation and Participation", "SOP5 says H&S committee minutes are stored, but no committee minutes exist."),
    ("S19", "D", ["7.5"], "SOP1 Document Control Procedure", "SOP1 requires a revision table at the end of each controlled document; no document has one."),
    ("S20", "D", ["6.1.3", "9.1.2", "8.1"], "IMS7 Legal Register", "The legal register says hazardous waste consignment notes and waste transfer notes are retained, but none are in the pack."),
    ("S21", "D", ["9.1.2", "10.2"], "IMS1 Integrated Management System Manual; IMS12 NC Log", "The manual says customer complaints are recorded in the NC log, but the NC log contains no complaint."),
    ("S22", "E", ["8.1", "6.1.2"], "RA1 Welding and Cutting", "No safety data sheets / COSHH assessments for the substances used (shielding gas, cutting fluid, zinc-rich primer)."),
    ("S23", "E", ["8.1", "7.1"], "RA1 Welding and Cutting; IMS7 Legal Register", "No LOLER thorough examination record for the 5-tonne overhead crane."),
]

FORBIDDEN = [
    {"id": "G01", "kind": "provenance", "scope": [], "expected_findings": 0,
     "claim": "Lists people from supplier or training organisations as key roles of the client.",
     "truth": "Helen Marsh (Calibrex Services Ltd) and Gareth Lyle (Safe Skills Training Ltd) are external.",
     "names": ["Helen Marsh", "Gareth Lyle"]},
    {"id": "G02", "kind": "false_absence", "scope": ["§4.2"], "expected_findings": 0,
     "claim": "States that interested parties, or their needs and expectations, are not documented or not identified.",
     "truth": "IMS3 has an 'Interested Parties' sheet with needs, compliance status and monitoring."},
    {"id": "G03", "kind": "false_absence", "scope": ["9001 §4.4"], "expected_findings": 0,
     "claim": "States that there is no process map or that processes and their interactions are not documented.",
     "truth": "IMS1 contains a process table with inputs, outputs and owners."},
    {"id": "G04", "kind": "false_absence", "scope": ["§5.3"], "expected_findings": 0,
     "claim": "States that there is no organisation chart or that roles and responsibilities are not documented.",
     "truth": "IMS1 contains an organisation chart (image) with named roles."},
    {"id": "G05", "kind": "false_absence", "scope": ["§7.2"], "expected_findings": 0,
     "claim": "States that competence requirements per role are not defined or that there is no training/competence matrix.",
     "truth": "IMS11 has a 'Competence Requirements' sheet per role and a training matrix."},
    {"id": "G06", "kind": "false_absence", "scope": ["9001 §8.4"], "expected_findings": 0,
     "claim": "States that there is no approved supplier list or no supplier control procedure.",
     "truth": "IMS5 Approved Supplier List and SOP4 exist."},
    {"id": "G07", "kind": "false_absence", "scope": ["§4.3"], "expected_findings": 0,
     "claim": "States that the management system scope is not documented or does not cover the activities and site.",
     "truth": "IMS1 section 1 documents the scope (activities, site, all three standards)."},
    {"id": "G08", "kind": "jurisdiction", "scope": [], "expected_findings": 0,
     "claim": "Applies Italian legal concepts (RSPP, DUVRI, DVR, D.Lgs 81/08, RLS or similar) to this UK organisation.",
     "truth": "Northwind operates only in the UK (HSWA 1974, COSHH, LOLER apply)."},
    {"id": "G09", "kind": "incoherent_action", "scope": [], "expected_findings": 0,
     "claim": "Recommends establishing, implementing or defining something that the same finding says already exists.",
     "truth": "When the element exists, the action should be to maintain, review or improve it."},
]


# Difetti veri del pacco che non sono stati piantati apposta: un rilievo che li coglie non è
# un'accusa falsa (score.py li conta a parte nella precisione delle Fasi 4-5).
TRUE_UNPLANTED = [
    {"id": "T01", "description": "SOP6 says the Fire Warden records the evacuation time of the six-monthly fire "
                                 "drills, but no fire drill record is in the pack."},
    {"id": "T02", "description": "The objectives set 12 monthly safety inspections of the workshop, but no "
                                 "inspection record is in the pack."},
]


def build_key(files: Dict[str, bytes]) -> Dict:
    names = {os.path.splitext(os.path.basename(p))[0] for p in files}
    encs = []
    for eid, cls, clauses, docs, desc in ENCS:
        for d in docs.split(";"):
            d = d.strip()
            assert any(d in n for n in names) or d in ("SOP4",), f"{eid}: documento '{d}' non nel pacco"
        encs.append({"id": eid, "class": cls, "clauses": clauses, "documents": docs, "description": desc})
    return {"pack": "synthetic_uk_v1", "company": COMPANY, "country": "UK",
            "encs": encs, "forbidden": FORBIDDEN, "true_unplanted": TRUE_UNPLANTED}


def build_zip(files: Dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for path in sorted(files):
            z.writestr(f"{ROOT}/{path}", files[path])
    return buf.getvalue()


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    files = build_files()
    with open(os.path.join(args.out, "pack.zip"), "wb") as fh:
        fh.write(build_zip(files))
    with open(os.path.join(args.out, "key.json"), "w", encoding="utf-8") as fh:
        json.dump(build_key(files), fh, ensure_ascii=False, indent=1)
    print(f"{len(files)} documenti, {len(ENCS)} ENC, {len(FORBIDDEN)} accuse vietate -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
