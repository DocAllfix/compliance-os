"""
build_registry.py — Piano motore Fase 2, ricostruzione della knowledge base,
passi 2-5: aggancio norma ↔ checklist ↔ KB attuale e tabella di revisione.

Per ogni riga della norma (da decompose_norm.py) il modello produce, con
parole nostre e in inglese (il testo della norma non si copia):
  - en:            il requisito riformulato (≤ 30 parole);
  - object:        record | document | process | outcome | behaviour;
  - verifiable:    high | medium | none (dai soli documenti);
  - kb_match:      requisiti atomici della KB attuale che verificano lo stesso obbligo;
  - checklist_how: come lo verifica un auditor nei documenti (dalla checklist).
I requisiti della KB attuale che non corrispondono a nessun "deve" vengono
classificati: good_practice (al massimo opportunità) | national_IT (profilo
Italia) | duplicate | other.
La gravità massima la calcola Python dagli attributi (mai il modello).

Uscite (FUORI dal repo: contengono testo della norma per la revisione):
  registry_<norma>.json, registry_review.xlsx

Uso:
  python tools/kb/build_registry.py --kb-dir C:/Users/user/amisys-bench/kb
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from typing import Any, Dict, List

_BACKEND = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, _BACKEND)

from dotenv import load_dotenv  # noqa: E402
load_dotenv(os.path.join(_BACKEND, ".env"))

from compliance.gap_analyzer import _NORM_REGISTRY  # noqa: E402

NORMS = {"9001": "ISO 9001:2015", "14001": "ISO 14001:2015", "45001": "ISO 45001:2018"}
CHECKLISTS = {"9001": "chk_Template9001.txt", "14001": "chk_TEMPLATE14001.txt", "45001": "chk_TEMPLATE450001.txt"}

# 14001: testo disponibile 2026, struttura supportata 2015 (decisione utente).
MAP_14001_2026_TO_2015 = {"6.1.4": "6.1.1", "6.1.5": "6.1.4", "9.3.1": "9.3", "9.3.2": "9.3", "9.3.3": "9.3",
                          "10.1": "10.3"}
ONLY_2026 = {"6.3"}

PROMPT = """You build the requirement registry of an ISO compliance engine for clause {clause} of {norm}.

NORM ROWS (official Italian text of the standard — never copy it, rewrite in your own English words):
{rows}

CURRENT KNOWLEDGE-BASE REQUIREMENTS for this clause:
{kb}

AUDIT CHECKLIST GUIDANCE for this clause (Italian):
{checklist}

For every norm row return:
- "en": the requirement in your own UK English words, max 30 words, stating what must be true;
- "object": "record" (evidence of results to retain), "document" (information to maintain/make available),
  "process" (a process/method must exist and be applied), "outcome" (a result/property), or "behaviour"
  (leadership, awareness, attitudes — not provable from documents);
- "verifiable": "high" (a document/record directly shows it), "medium" (indirect evidence), "none";
- "kb_match": ids of current knowledge-base requirements that check the same obligation (may be empty);
- "checklist_how": max 25 words, how an auditor checks it in documents (use the checklist guidance if relevant).
For every current knowledge-base requirement that matches NO norm row return its class:
"good_practice" (useful but not required by the standard), "national_IT" (Italian legal obligation),
"duplicate", or "other", with a short reason.

Return JSON: {{"rows": {{"<row_id>": {{"en": "...", "object": "...", "verifiable": "...", "kb_match": [...],
"checklist_how": "..."}}}}, "kb_unmatched": {{"<kb id>": {{"class": "...", "reason": "..."}}}}}}"""


def parse_checklist(path: str) -> Dict[str, str]:
    lines = [l.strip() for l in open(path, encoding="utf-8").read().splitlines()]
    out: Dict[str, List[str]] = {}
    cur = None
    for i, l in enumerate(lines):
        if re.fullmatch(r"\d{1,2}(\.\d{1,2}){0,3}", l):
            cur = l
            out.setdefault(cur, [])
            continue
        if cur and l and not l.startswith("<w:") and not l.startswith("{{") and "w:val=" not in l:
            out[cur].append(l)
    return {k: "\n".join(v) for k, v in out.items()}


def kb_clause_for(norm: str, clause: str, kb_ids: List[str]) -> str:
    """Clausola della KB (id 9001_8_4) che contiene la clausola di testo (8.4.1)."""
    if norm == "14001":
        clause = MAP_14001_2026_TO_2015.get(clause, clause)
    parts = clause.split(".")
    for k in range(len(parts), 0, -1):
        cid = f"{norm}_" + "_".join(parts[:k])
        if cid in kb_ids:
            return cid
    return f"{norm}_" + "_".join(parts[:2])          # clausola nuova, non ancora nella KB


def max_severity(row: Dict[str, Any]) -> str:
    """Tetto di gravità dagli attributi del testo (Principio 2 del piano)."""
    if row["strength"] in ("should", "may", "note"):
        return "OBSERVATION"
    if "consider" in row["verbs"] and "take_into_account" not in row["verbs"]:
        return "OBSERVATION"
    if row.get("object") == "behaviour" or row.get("verifiable") == "none":
        return "TO_VERIFY"
    return "MINOR"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--kb-dir", required=True)
    ap.add_argument("--norms", nargs="+", default=list(NORMS))
    args = ap.parse_args()
    from v2.azure_openai_client_v2 import make_client
    client = make_client(deep=True)

    def ask(prompt: str) -> Dict[str, Any]:
        for _ in range(3):
            try:
                r = client._sdk.chat.completions.create(
                    model=client.deployment, temperature=0.0, seed=42, max_tokens=8000,
                    response_format={"type": "json_object"}, messages=[{"role": "user", "content": prompt}])
                return json.loads(r.choices[0].message.content)
            except Exception as e:
                print(f"   ritento: {str(e)[:120]}")
        return {}

    review_rows, unmatched_rows = [], []
    for norm in args.norms:
        rows = json.load(open(os.path.join(args.kb_dir, f"rows_{norm}.json"), encoding="utf-8"))
        checklist = parse_checklist(os.path.join(args.kb_dir, CHECKLISTS[norm]))
        kb = _NORM_REGISTRY[NORMS[norm]]["clauses"]
        groups: Dict[str, List[Dict]] = {}
        for r in rows:
            if norm == "14001" and r["clause"] in ONLY_2026:
                continue
            groups.setdefault(kb_clause_for(norm, r["clause"], list(kb)), []).append(r)

        def work(item):
            cid, rs = item
            clause_kb = kb.get(cid, {})
            kb_reqs = clause_kb.get("atomic_requirements") or []
            num = ".".join(cid.split("_")[1:])
            chk = "\n".join(v for k, v in checklist.items() if k == num or k.startswith(num + "."))[:6000]
            prompt = PROMPT.format(
                clause=num, norm=NORMS[norm],
                rows="\n".join(f"- {r['row_id']} [{r['strength']}] {r['text']}" for r in rs),
                kb="\n".join(f"- {q['id']}: {q['text']} (risk_if_missing {q.get('risk_if_missing')})"
                             for q in kb_reqs) or "(none)",
                checklist=chk or "(none)")
            return cid, rs, kb_reqs, ask(prompt)

        with ThreadPoolExecutor(max_workers=6) as ex:
            results = list(ex.map(work, sorted(groups.items())))
        registry = []
        for cid, rs, kb_reqs, out in results:
            got = out.get("rows") or {}
            for r in rs:
                g = got.get(r["row_id"]) or {}
                entry = {**r, "kb_clause": cid, "en": g.get("en", ""), "object": g.get("object", r["object"]),
                         "verifiable": g.get("verifiable", "medium"),
                         "kb_match": [m for m in g.get("kb_match") or [] if any(m == q["id"] for q in kb_reqs)],
                         "checklist_how": g.get("checklist_how", "")}
                entry["max_severity"] = max_severity(entry)
                registry.append(entry)
                review_rows.append(entry)
            for kid, v in (out.get("kb_unmatched") or {}).items():
                q = next((q for q in kb_reqs if q["id"] == kid), None)
                if q:
                    unmatched_rows.append({"norm": norm, "kb_clause": cid, "id": kid, "text": q["text"],
                                           "risk_if_missing": q.get("risk_if_missing"),
                                           "class": v.get("class"), "reason": v.get("reason")})
        with open(os.path.join(args.kb_dir, f"registry_{norm}.json"), "w", encoding="utf-8") as f:
            json.dump(registry, f, ensure_ascii=False, indent=1)
        print(f"{norm}: {len(registry)} righe di norma, {sum(1 for u in unmatched_rows if u['norm'] == norm)} "
              f"requisiti KB senza obbligo dietro")

    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Norm requirements"
    head = ["Norm", "KB clause", "Text clause", "Row", "Strength", "Verbs", "Object", "Verifiable",
            "Max severity", "Requirement (EN, our words)", "Official text (IT, local review only)",
            "Current KB match", "Checklist: how to verify", "Reviewer decision", "Reviewer note"]
    ws.append(head)
    for r in review_rows:
        ws.append([r["norm"], r["kb_clause"], r["clause"], r["row_id"], r["strength"], ", ".join(r["verbs"]),
                   r["object"], r["verifiable"], r["max_severity"], r["en"], r["text"], ", ".join(r["kb_match"]),
                   r["checklist_how"], "", ""])
    ws2 = wb.create_sheet("KB without norm basis")
    ws2.append(["Norm", "KB clause", "KB requirement id", "Text", "Risk if missing (today)", "Class", "Reason",
                "Proposed fate", "Reviewer decision"])
    fate = {"good_practice": "Keep as good practice: max opportunity for improvement, never NC",
            "national_IT": "Move to Italy profile (only for Italian organisations)",
            "duplicate": "Drop (covered by another requirement)", "other": "Review"}
    for u in unmatched_rows:
        ws2.append([u["norm"], u["kb_clause"], u["id"], u["text"], u["risk_if_missing"], u["class"], u["reason"],
                    fate.get(u["class"], "Review"), ""])
    for sh in (ws, ws2):
        for c in sh[1]:
            c.font = Font(bold=True)
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width = 18
        sh.freeze_panes = "A2"
    for col in ("J", "K", "M"):
        ws.column_dimensions[col].width = 60
        for c in ws[col]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["D"].width = 60
    ws2.column_dimensions["G"].width = 50
    out = os.path.join(args.kb_dir, "registry_review.xlsx")
    wb.save(out)
    print(f"tabella di revisione: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
