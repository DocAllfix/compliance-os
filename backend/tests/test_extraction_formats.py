"""
Fast regression — Fase A (coverage estensioni).

Per ogni handler nuovo: campione sintetico → `extract_document` produce testo
plausibile. I formati storici (docx/xlsx) restano INVARIATI (asserito). Gli
handler non sollevano mai: su input non valido tornano "".

I formati non sintetizzabili offline in modo affidabile (.doc OLE reale, .msg
reale, immagini→OCR) sono coperti da: (a) test di robustezza "non solleva, torna
''"; (b) test di ROUTING (la dispatch instrada l'estensione all'handler giusto).
"""
from __future__ import annotations

import io

import compliance.extraction as extraction
from compliance.extraction import extract_document


# ── Formati con campione reale sintetizzabile ────────────────────────────────

def test_docx_unchanged():
    import docx
    d = docx.Document()
    d.add_paragraph("Hello DOCX paragraph")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "ColA"
    t.rows[0].cells[1].text = "ColB"
    buf = io.BytesIO(); d.save(buf)
    out = extract_document("x.docx", buf.getvalue())
    assert "Hello DOCX paragraph" in out
    assert "ColA | ColB" in out          # tabelle (comportamento storico)


def test_xlsx_unchanged():
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "S1"
    ws["A1"] = "KPI"; ws["B1"] = "Target"; ws["A2"] = "uptime"; ws["B2"] = 99
    buf = io.BytesIO(); wb.save(buf)
    out = extract_document("x.xlsx", buf.getvalue())
    assert '[SHEET "S1"]' in out
    assert "KPI | Target" in out


def test_pptx():
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "Hello PPTX slide"
    buf = io.BytesIO(); prs.save(buf)
    out = extract_document("deck.pptx", buf.getvalue())
    assert "Hello PPTX slide" in out


def test_odt():
    from odf.opendocument import OpenDocumentText
    from odf.text import P
    doc = OpenDocumentText()
    doc.text.addElement(P(text="Hello ODT body"))
    buf = io.BytesIO(); doc.save(buf)
    out = extract_document("doc.odt", buf.getvalue())
    assert "Hello ODT body" in out


def test_ods():
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableRow, TableCell
    from odf.text import P
    doc = OpenDocumentSpreadsheet()
    tbl = Table(name="Sheet1")
    tr = TableRow()
    cell = TableCell(); cell.addElement(P(text="Hello ODS cell"))
    tr.addElement(cell); tbl.addElement(tr)
    doc.spreadsheet.addElement(tbl)
    buf = io.BytesIO(); doc.save(buf)
    out = extract_document("sheet.ods", buf.getvalue())
    assert "Hello ODS cell" in out


def test_eml():
    from email.message import EmailMessage
    m = EmailMessage()
    m["From"] = "a@b.com"; m["To"] = "c@d.com"; m["Subject"] = "Test Subject"
    m.set_content("Hello EML body text")
    out = extract_document("mail.eml", m.as_bytes())
    assert "Subject: Test Subject" in out
    assert "Hello EML body text" in out


def test_html_strips_scripts():
    raw = b"<html><body><h1>Hello HTML</h1><script>evil()</script></body></html>"
    out = extract_document("page.html", raw)
    assert "Hello HTML" in out
    assert "evil()" not in out


def test_xml():
    raw = b"<root><a>Hello XML</a><b>second</b></root>"
    out = extract_document("data.xml", raw)
    assert "Hello XML" in out and "second" in out


def test_rtf():
    raw = br"{\rtf1\ansi\deff0 Hello RTF text}"
    out = extract_document("note.rtf", raw)
    assert "Hello RTF text" in out


def test_p7m_unwraps_embedded_docx():
    # Un .p7m che incapsula un DOCX: prefisso "firma" + bytes DOCX reali.
    import docx
    d = docx.Document(); d.add_paragraph("Inside the signed envelope")
    buf = io.BytesIO(); d.save(buf)
    p7m_bytes = b"\x30\x82\x01\x00GARBAGE-PKCS7-HEADER" + buf.getvalue()
    out = extract_document("contract.pdf.p7m", p7m_bytes)
    assert "Inside the signed envelope" in out


# ── Robustezza: handler non sollevano mai, tornano "" su input invalido ──────

def test_doc_legacy_graceful_on_garbage():
    assert extract_document("x.doc", b"not an OLE file at all") == ""


def test_msg_graceful_on_garbage():
    assert extract_document("x.msg", b"not a real msg") == ""


def test_unknown_extension_returns_empty():
    assert extract_document("x.weirdext", b"whatever") == ""


# ── Routing: l'estensione immagine è instradata all'OCR (senza rete) ─────────

def test_image_routes_to_ocr(monkeypatch):
    called = {}

    def fake_ocr(raw, filename=""):
        called["hit"] = filename
        return "OCR-TEXT-SENTINEL"

    monkeypatch.setattr(extraction, "_extract_pdf_text", extraction._extract_pdf_text)
    import compliance.azure_ocr as azure_ocr
    monkeypatch.setattr(azure_ocr, "ocr_image_bytes", fake_ocr)

    out = extract_document("scan.png", b"\x89PNG\r\n")
    assert out == "OCR-TEXT-SENTINEL"
    assert called.get("hit") == "scan.png"
