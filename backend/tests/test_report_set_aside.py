"""
Fast regression for the set-aside filter in the DOCX report (Fase 4/8).

Pins the leak found in the Fase 8 e2e: a dismissed finding still appeared
in the per-norm "Top Priority Findings" list of the Executive Summary.
Dismissed clauses must vanish from every operative section and appear only
in the Clause Matrix (marked "Set aside") and the Set-aside appendix.
"""
from __future__ import annotations

import io

from docx import Document

from compliance.report_builder import build_docx_report
from compliance.schemas import (
    ClauseGap,
    CompanyContext,
    ComplianceAnalysisResult,
    NormGapReport,
    RemediationPlan,
    TraceabilityLink,
)


def _gap(clause_id: str, title: str, severity: str, score: int) -> ClauseGap:
    return ClauseGap(
        clause_id=clause_id,
        clause_title=title,
        score=score,
        severity=severity,
        is_spof=False,
        finding=f"Finding for {title}.",
        sources=[TraceabilityLink(
            evidence_id="ev_001", file_name="doc.docx", pages=[],
            excerpt="…", relevance_score=0.5,
        )],
        remediation=RemediationPlan(
            what="Fix it", who="Quality Manager", when="30 days", effort="LOW",
        ),
    )


def _result() -> ComplianceAnalysisResult:
    gaps = [
        _gap("9001_7_5", "7.5 Documented information", "MAJOR", 20),
        _gap("9001_6_2", "6.2 Objectives", "MINOR", 55),
        _gap("9001_9_3", "9.3 Management review", "MAJOR", 30),
    ]
    return ComplianceAnalysisResult(
        analysis_id="an_test",
        company="TestCo",
        context=CompanyContext(
            company_name="TestCo", sector="Engineering", employee_range="11-50",
            certified="yes", sites_count="1", norms=["ISO 9001:2015"],
        ),
        analysis_date="2026-06-12",
        norms=[NormGapReport(
            norm="ISO 9001:2015", overall_score=40, clauses=gaps,
            major_count=2, minor_count=1, observation_count=0,
            conformant_count=0, spof_count=0,
        )],
        documents_analyzed=3,
        top_gaps=gaps[:1],
    )


def _all_text(docx_bytes: bytes) -> tuple[list[str], list[list[str]]]:
    doc = Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    rows = [[c.text for c in row.cells] for tbl in doc.tables for row in tbl.rows]
    return paragraphs, rows


def test_dismissed_clause_leaves_operative_sections():
    register = [{
        "id": 1, "analysis_id": "an_test", "clause_id": "9001_7_5",
        "source": "engine", "title": "7.5 Documented information",
        "severity": "MAJOR", "dismissed": 1,
        "dismiss_reason": "Handled in the customer portal.", "created_by": "alice",
    }]
    paragraphs, rows = _all_text(build_docx_report(_result(), register=register))

    # niente 7.5 in nessun paragrafo operativo (Top Priority, Critical, Action Plan)
    leaks = [p for p in paragraphs if "7.5 Documented information" in p]
    assert leaks == [], f"dismissed finding leaked into operative sections: {leaks}"

    # in Clause Matrix la riga resta ma marcata Set aside
    matrix_row = next(r for r in rows if r and r[0] == "9001_7_5")
    assert matrix_row[-1] == "Set aside"

    # appendice presente con motivazione
    assert any("SET-ASIDE" in p.upper() for p in paragraphs)
    assert any("Handled in the customer portal." in c for r in rows for c in r)


def test_without_dismissals_report_is_unchanged():
    paragraphs, rows = _all_text(build_docx_report(_result(), register=[]))
    assert any("7.5 Documented information" in p for p in paragraphs)
    assert not any("SET-ASIDE" in p.upper() for p in paragraphs)
    matrix_row = next(r for r in rows if r and r[0] == "9001_7_5")
    assert matrix_row[-1] == "MAJOR"
