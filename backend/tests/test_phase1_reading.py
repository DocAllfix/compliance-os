"""
Fast — Piano motore Fase 1 ("leggere tutto"), interruttori spenti di default.

Pin:
  - READ_SHEET_CELLS: celle vuote visibili, valori con la loro intestazione,
    intestazioni a due righe composte, date senza orario;
  - READ_SHEET_SAMPLING: ogni foglio di una cartella lunga arriva alla scheda e
    allo Stage B (non solo testa e coda);
  - READ_DOCX_IMAGES: le immagini del corpo di un Word passano dall'OCR al loro
    posto; icone e doppioni no;
  - READ_DOCX_HEADERS: il testo di intestazioni e piè di pagina (codice,
    versione, data) entra una volta sola in testa al documento;
  - RETRIEVE_PER_REQUIREMENT: una query per requisito atomico, risultati fusi
    senza doppioni;
  - a interruttore spento il comportamento è quello storico.
"""
from __future__ import annotations

import datetime as dt
import io

import pytest

from compliance import extraction as E
from compliance import sampling


@pytest.fixture(autouse=True)
def _all_off(monkeypatch):
    for v in ("ENGINE_PHASE1", "READ_SHEET_CELLS", "READ_SHEET_SAMPLING", "READ_DOCX_IMAGES",
              "READ_DOCX_HEADERS", "RETRIEVE_PER_REQUIREMENT"):
        monkeypatch.delenv(v, raising=False)


def _xlsx(sheets):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── READ_SHEET_CELLS ─────────────────────────────────────────────────────────

def test_blank_cells_keep_their_column():
    lines = E.render_sheet_rows([["Ref", "Risk", "Owner"], ["R1", "Crane", None], ["R2", None, "QM"]])
    assert "Row 2: Ref: R1 | Risk: Crane | Owner: (blank)" in lines
    assert "Row 3: Ref: R2 | Risk: (blank) | Owner: QM" in lines


def test_two_row_grouped_header_is_composed():
    lines = E.render_sheet_rows([
        [None, "First Aid", None, "LOTO", None],
        ["Name", "Date", "Expiry", "Date", "Expiry"],
        ["Ann", dt.datetime(2025, 2, 10), dt.datetime(2028, 2, 10), None, None],
    ])
    assert "Columns: Name | First Aid Date | First Aid Expiry | LOTO Date | LOTO Expiry" in lines
    assert any("First Aid Date: 2025-02-10" in l and "LOTO Date: (blank)" in l for l in lines)


def test_text_data_rows_are_not_mistaken_for_a_header():
    lines = E.render_sheet_rows([["Issue", "Type", "Effect"], ["Price", "External", "Margins"]])
    assert lines[0] == "Columns: Issue | Type | Effect"
    assert lines[1].startswith("Row 2: Issue: Price")


def test_flag_switches_spreadsheet_rendering(monkeypatch):
    raw = _xlsx({"S": [["A", "B", "C"], ["1", None, "3"]]})
    old = E.extract_document("x.xlsx", raw)
    assert "(blank)" not in old and "1 | 3" in old
    monkeypatch.setenv("READ_SHEET_CELLS", "1")
    new = E.extract_document("x.xlsx", raw)
    assert new.startswith('[SHEET "S"]') and "B: (blank)" in new


# ── READ_SHEET_SAMPLING ──────────────────────────────────────────────────────

def _workbook_text(n_sheets=4, rows=400):
    return "\n\n".join(f'[SHEET "Sheet{k}"]\n' + "\n".join(f"Row {i}: value {k}-{i}" for i in range(rows))
                       for k in range(n_sheets))


def test_sampling_off_is_head_tail():
    text = _workbook_text()
    assert sampling.for_card(text, 2000) == sampling.head_tail(text, 2000)
    assert "Sheet1" not in sampling.for_card(text, 2000)      # il foglio centrale si perde


def test_sampling_on_reaches_every_sheet(monkeypatch):
    monkeypatch.setenv("READ_SHEET_SAMPLING", "1")
    text = _workbook_text()
    card = sampling.for_card(text, 4000)
    assert all(f'[SHEET "Sheet{k}"]' in card for k in range(4))
    assert card.startswith("[WORKBOOK with 4 sheets")
    assert "more lines of this sheet omitted" in card
    assert len(card) <= 4000 + 400
    deep = sampling.for_deep(text, 4000)
    assert all(f'[SHEET "Sheet{k}"]' in deep for k in range(4))


def test_small_sheets_enter_whole(monkeypatch):
    monkeypatch.setenv("READ_SHEET_SAMPLING", "1")
    text = '[SHEET "Big"]\n' + "\n".join(f"r{i}" for i in range(3000)) + '\n\n[SHEET "Parties"]\nCustomers | delivery\n'
    s = sampling.sample_by_sheet(text, 3000)
    assert "Customers | delivery" in s


# ── READ_DOCX_IMAGES ─────────────────────────────────────────────────────────

def _docx_with_images(images):
    import docx
    from docx.shared import Inches
    d = docx.Document()
    d.add_paragraph("Section 5.3 Roles")
    for img in images:
        d.add_picture(io.BytesIO(img), width=Inches(2))
    d.add_paragraph("After the chart")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _png(size_px, seed=0):
    from PIL import Image
    import random
    rnd = random.Random(seed)
    img = Image.new("RGB", (size_px, size_px))
    img.putdata([(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)) for _ in range(size_px * size_px)])
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def test_docx_images_are_ocred_in_place(monkeypatch):
    calls = []
    monkeypatch.setattr("compliance.azure_ocr.ocr_image_bytes",
                        lambda raw, filename="": calls.append(filename) or "Managing Director")
    big, icon = _png(120, 1), _png(4, 2)
    raw = _docx_with_images([big, big, icon])      # doppione + icona
    assert "Managing Director" not in E.extract_document("m.docx", raw) and not calls
    monkeypatch.setenv("READ_DOCX_IMAGES", "1")
    text = E.extract_document("m.docx", raw)
    assert len(calls) == 1, "un'immagine letta una volta sola, l'icona ignorata"
    assert text.index("Section 5.3") < text.index("Managing Director") < text.index("After the chart")


# ── READ_DOCX_HEADERS ────────────────────────────────────────────────────────

def test_docx_header_and_footer_text_is_read_once(monkeypatch):
    import docx
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "QHSE Policy - Version 1 - 10/1/26"
    d.sections[0].footer.paragraphs[0].text = "Approved by: Managing Director"
    d.add_paragraph("We are committed to quality.")
    d.add_section()                                  # seconda sezione: stessa intestazione
    buf = io.BytesIO()
    d.save(buf)
    raw = buf.getvalue()
    assert "Version 1" not in E.extract_document("p.docx", raw)
    monkeypatch.setenv("READ_DOCX_HEADERS", "1")
    text = E.extract_document("p.docx", raw)
    assert text.startswith("[HEADER] QHSE Policy - Version 1 - 10/1/26")
    assert text.count("[HEADER]") == 1 and "[FOOTER] Approved by: Managing Director" in text
    assert text.index("[FOOTER]") < text.index("We are committed")


# ── RETRIEVE_PER_REQUIREMENT ─────────────────────────────────────────────────

class _FakeService:
    def __init__(self):
        self.queries = []

    def get_evidences_for_clause(self, clause, n=4):
        return [{"evidence_id": "ev_001", "excerpt": "policy text", "distance": 0.3}]

    def get_evidences_similar_to(self, q, n=6):
        self.queries.append(q)
        return [{"evidence_id": f"ev_{len(self.queries) + 10:03d}", "excerpt": q[:40], "distance": 0.5}]


def test_per_requirement_queries(monkeypatch):
    from compliance.gap_analyzer import _retrieve_with_fallback
    clause = {"id": "9001_9_3", "title": "9.3", "evidence_keywords": [],
              "atomic_requirements": [{"text": "inputs a"}, {"text": "inputs b", "verification_hint": "minutes"}]}
    svc = _FakeService()
    base = _retrieve_with_fallback(svc, clause)
    assert not svc.queries and [r["evidence_id"] for r in base] == ["ev_001"]
    monkeypatch.setenv("RETRIEVE_PER_REQUIREMENT", "1")
    out = _retrieve_with_fallback(svc, clause)
    assert svc.queries == ["inputs a", "inputs b minutes"]
    assert [r["evidence_id"] for r in out][0] == "ev_001" and len(out) == 3
